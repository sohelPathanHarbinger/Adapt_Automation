"""Word styles and helpers shared by the storyboard generators.

A generated storyboard is written with the exact styles the parser reads, so
the document is not a description of the conventions - it is a course that
builds. Both `make_base_storyboard.py` and the Template Reference generator
use this module, so the two documents can never drift apart stylistically.
"""

from __future__ import annotations

from pathlib import Path

from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor

#: name, size, bold, colour. The colours are for the human reading the .docx;
#: the parser only ever looks at the style name.
PARAGRAPH_STYLES = [
    ("Interactivity-Heading-1", 12, True, "1F4E79"),
    ("Interactivity-Heading-2", 11, True, "2E74B5"),
    ("Interactivity-body", 11, False, None),
    ("Interactivity-Bullet-1", 11, False, None),
    ("Interactive-bullet-2", 10, False, None),
    ("Interactivity-Label", 9, False, "595959"),
    ("CYP-Question", 11, True, None),
    ("CYP - Answer", 11, False, None),
    ("CYP - TrueFalse", 11, False, None),
    ("Bullet list level 1", 11, False, None),
    ("Bullet list level 2", 10, False, None),
    ("Number bullet list 1", 11, False, None),
    ("Diagram label 1", 9, True, "595959"),
    ("footnote/diagram label", 9, True, "595959"),
    ("il Reference", 9, False, "595959"),
    ("Refrence text", 9, False, "C00000"),
    ("Programming-Note", 10, False, "7030A0"),
]

CHARACTER_STYLES = [
    ("Glossary item Char", "1F7A3D", True),
    ("Refrence text Char", "C00000", False),
    ("Programming-Notes", "7030A0", False),
]


def ensure_styles(doc) -> None:
    """Add every style the parser reads, so an author can copy the file."""
    styles = doc.styles
    existing = {s.name for s in styles}

    for name, size, bold, colour in PARAGRAPH_STYLES:
        if name in existing:
            continue
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles["Normal"]
        style.font.size = Pt(size)
        style.font.bold = bold
        if colour:
            style.font.color.rgb = RGBColor.from_string(colour)

    for name, colour, bold in CHARACTER_STYLES:
        if name in existing:
            continue
        style = styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
        style.font.color.rgb = RGBColor.from_string(colour)
        style.font.bold = bold


def p(doc, style: str, text: str = ""):
    """One paragraph in a named style."""
    para = doc.add_paragraph(style=style)
    if text:
        para.add_run(text)
    return para


def runs(doc, style: str, parts):
    """A paragraph built from ``(text, character-style-or-None)`` pairs."""
    para = doc.add_paragraph(style=style)
    for text, char in parts:
        run = para.add_run(text)
        if char:
            run.style = doc.styles[char]
    return para


def table(doc, rows, header: bool = True):
    """Build a table.

    A cell given as a list holds one paragraph per entry, each optionally
    ``(style, text)`` - which is how a Key Concepts panel carries its topic
    line and then its bullets.
    """
    grid = doc.add_table(rows=len(rows), cols=len(rows[0]))
    grid.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            if isinstance(cell, (list, tuple)):
                target = grid.rows[r].cells[c]
                for n, entry in enumerate(cell):
                    style, text = entry if isinstance(entry, tuple) else (None, entry)
                    para = target.paragraphs[0] if n == 0 else target.add_paragraph()
                    para.text = str(text)
                    if style:
                        para.style = doc.styles[style]
                continue
            grid.rows[r].cells[c].text = str(cell)
            if r == 0 and header:
                for para in grid.rows[r].cells[c].paragraphs:
                    for run in para.runs:
                        run.bold = True
    doc.add_paragraph()
    return grid


def box(doc, lines):
    """One interactivity drawn inside a 1x1 table, as most storyboards do.

    ``lines`` is a list of ``(style, text)`` pairs, or ``("image", path,
    width_inches)`` for a picture inside the box.
    """
    from docx.shared import Inches

    grid = doc.add_table(rows=1, cols=1)
    grid.style = "Table Grid"
    cell = grid.rows[0].cells[0]
    first = True

    def next_para(style=None):
        nonlocal first
        para = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        if style:
            para.style = doc.styles[style]
        return para

    for line in lines:
        if line[0] == "image":
            para = next_para()
            para.add_run().add_picture(str(line[1]), width=Inches(line[2]))
            continue
        style, text = line[0], line[1]
        para = next_para(style)
        if isinstance(text, list):
            for chunk, char in text:
                run = para.add_run(chunk)
                if char:
                    run.style = doc.styles[char]
        else:
            para.add_run(text)
    doc.add_paragraph()
    return grid


def picture(doc, path: Path, inches: float = 4.5):
    """A figure on its own paragraph."""
    from docx.shared import Inches

    doc.add_picture(str(path), width=Inches(inches))
    return doc.paragraphs[-1]
