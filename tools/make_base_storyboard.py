"""Generate `base/Inputs`'s storyboard - the one the base course is built from.

`base/Source` holds one sample of every component the automation can emit, and
this document is the storyboard that produces it. Keeping the two in step is
the point: change a component here, rebuild, and `base/Source` follows.

    python tools/make_base_storyboard.py          # write the .docx
    python build.py base/Inputs --apply           # rebuild base/Source from it

Every picture, video and PDF it uses comes from `base/Inputs/assets`, so the
base course ships a handful of media files rather than a course's worth.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))

from storyboard_docx import box, ensure_styles, p, picture, runs, table  # noqa: E402

DEST = ROOT / "base" / "Inputs"
MEDIA = DEST / "assets" / "images" / "Course"
BRIGHT = MEDIA / "bright_image.jpg"
DARK = MEDIA / "dark_image.jpg"
AVATAR = MEDIA / "avatar.png"
ICON = MEDIA / "icon-avatar.png"

TITLE = "Base Sample Module"


def intro(doc) -> None:
    p(doc, "Programming-Note",
      "Programming Note: This 'How to Use This Module' page is for Print only. "
      "The programmed Adapt 'How to Use This Module' page should mirror the "
      "Figma template.")
    doc.add_heading("How to Use This Module", 1)
    runs(doc, "Normal", [
        ("This module is divided into chapters and sections. Each chapter begins "
         "with a list of Learning Objectives (LOs). Medical terminology is ", None),
        ("highlighted in boldface type", "Glossary item Char"),
        (" and defined in the Glossary at the end of the module.", None),
    ])

    doc.add_heading("Introduction", 1)
    doc.add_heading("Welcome", 2)
    runs(doc, "Normal", [
        ("Welcome to the ", None),
        ("base sample module", "Glossary item Char"),
        (". This module is not about a product: it exists so that every "
         "component the automation can build appears once, with real copy "
         "around it. ", None),
        ("[Base Source 2026 p1/A]", "Refrence text Char"),
        (" Use it to see what a component looks like before authoring your own "
         "storyboard.", None),
    ])

    doc.add_heading("Module Summary", 2)
    p(doc, "Normal", "This module is arranged as a course, not as a catalogue:")
    p(doc, "Bullet list level 1",
      "Chapter 1 covers the content components: copy, lists, tables, figures "
      "and callouts.")
    p(doc, "Bullet list level 1",
      "Chapter 2 covers the interactivities: accordions, tabs, hot spots, "
      "carousels, popups and sliders.")
    p(doc, "Bullet list level 1",
      "Each chapter ends with Key Concepts and a set of Check Your Progress "
      "questions, as a real module does.")

    doc.add_heading("Module Learning Objectives", 2)
    p(doc, "Normal", "Upon completion of this module, learners will be able to:")
    p(doc, "Bullet list level 1", "Recognise every component the automation builds.")
    p(doc, "Bullet list level 1", "Author a storyboard that asks for each of them.")
    p(doc, "Bullet list level 1", "Tell the conventions that carry meaning from those that do not.")


def chapter_one(doc) -> None:
    doc.add_heading("Chapter 1: Content Components", 1)
    box(doc, [
        ("Programming-Note", "Interactivity: Accordion"),
        ("Heading 4", "Learning Objectives"),
        ("Normal", "Upon completion of this chapter, learners will be able to:"),
        ("Bullet list level 1", "Identify how body copy, lists and tables are authored."),
        ("Bullet list level 1", "Describe how a figure and its caption travel together."),
        ("Bullet list level 1", "Recognise the three callout types."),
    ])

    doc.add_heading("Section 1.1: Body Copy, Lists and Tables", 2)
    doc.add_heading("Body copy", 3)
    runs(doc, "Normal", [
        ("A paragraph becomes a text component. A ", None),
        ("glossary term", "Glossary item Char"),
        (" becomes a popup whose definition comes from the Glossary table, and "
         "a citation is removed before the learner sees it. ", None),
        ("[Base Source 2026 p2/A]", "Refrence text Char"),
        (" Trademarks such as Adapt", None),
    ])
    p(doc, "Bullet list level 1", "A first-level bullet.")
    p(doc, "Bullet list level 2", "A second-level bullet, nested inside it.")
    p(doc, "Normal", "A numbered list reads as an ordered list:")
    p(doc, "Number bullet list 1", "The first step.")
    p(doc, "Number bullet list 1", "The second step.")

    doc.add_heading("A list with icons", 3)
    p(doc, "Normal", "A small picture inside a list item becomes that item's bullet:")
    for text in ("Each item carries its own icon.",
                 "The icon is any picture no wider than 200 pixels.",
                 "Larger pictures are placed after the list instead."):
        para = doc.add_paragraph(style="Bullet list level 1")
        para.add_run().add_picture(str(ICON), width=Inches(0.3))
        para.add_run(" " + text)

    doc.add_heading("A data table", 3)
    p(doc, "Normal", "The caption below names the table, so it heads it.")
    p(doc, "Diagram label 1", "Table 1-1: A Data Table")
    table(doc, [
        ["Parameter", "Group A", "Group B"],
        ["Mean at baseline", "4.28 (0.16)", "4.57 (0.23)"],
        ["Mean at week 52", "5.75 (0.15)", "3.95 (0.14)"],
    ])

    doc.add_heading("Section 1.2: Figures and Callouts", 2)
    doc.add_heading("A figure and its caption", 3)
    p(doc, "Normal",
      "A picture becomes a graphic component, and the caption under it becomes "
      "the figure's attribution. The learner can enlarge it.")
    picture(doc, BRIGHT, 4.5)
    p(doc, "Diagram label 1", "Figure 1-1: A Sample Figure")
    box(doc, [
        ("Programming-Note", "Interactivity: Callout"),
        ("Heading 4", "ZOOM IN"),
        ("Normal",
         "A callout attaches to the component above it - here, to the figure. "
         "ZOOM IN, QUICK FACT! and FAST FORWARD / FLASH BACK each pick their "
         "own icon."),
    ])

    doc.add_heading("Two callouts in a row", 3)
    p(doc, "Normal",
      "Where a second callout follows with no copy between, it is given a "
      "component of its own to sit on.")
    box(doc, [
        ("Programming-Note", "Interactivity: Callout"),
        ("Heading 4", "QUICK FACT!"),
        ("Normal", "A quick fact points out something relevant to the copy beside it."),
    ])
    box(doc, [
        ("Programming-Note", "Interactivity: Callout"),
        ("Heading 4", "FAST FORWARD / FLASH BACK"),
        ("Normal", "This callout points backwards or forwards to another section."),
    ])

    doc.add_heading("Section 1.3: Side by Side", 2)
    p(doc, "Programming-Note",
      "Programming Note: Display image side by side with paragraph below.")
    picture(doc, AVATAR, 1.6)
    p(doc, "Normal",
      "A programming note that asks for a side-by-side layout is carried out: "
      "the picture and this paragraph share the row, the picture on the left "
      "because it comes first.")

    doc.add_heading("Key Concepts (For Print)", 2)
    p(doc, "Normal",
      "This copy is written for the print PDF and is not built into the course.")

    doc.add_heading("Key Concepts (For Programming)", 2)
    table(doc, [
        ["Narration for KCs Presented by Avatar", "On Screen Text"],
        ["Copy is authored as body text, lists and tables, and a figure carries "
         "its caption with it.",
         ["Content Components",
          ("Bullet list level 1", "Body copy, bullets and numbered lists"),
          ("Bullet list level 1", "Tables, figures and captions")]],
        ["Callouts attach to the component above them, and each type has its "
         "own icon.",
         ["Callouts",
          ("Bullet list level 1", "ZOOM IN, QUICK FACT!, FAST FORWARD / FLASH BACK"),
          ("Bullet list level 1", "A second callout gets its own component")]],
    ])

    doc.add_heading("Check Your Progress", 2)
    p(doc, "CYP-Question", "Which component does a plain paragraph become?")
    para = doc.add_paragraph(style="CYP - Answer")
    para.add_run("A text component").bold = True
    p(doc, "CYP - Answer", "A graphic component")
    p(doc, "CYP - Answer", "An accordion")
    p(doc, "CYP - Answer", "A callout")

    p(doc, "CYP-Question",
      "Which of the following travel with the figure above them? Select all that apply.")
    para = doc.add_paragraph(style="CYP - Answer")
    para.add_run("Its caption").bold = True
    para = doc.add_paragraph(style="CYP - Answer")
    para.add_run("A callout that follows it").bold = True
    p(doc, "CYP - Answer", "The page title")
    para = doc.add_paragraph(style="CYP - Answer")
    para.add_run("Its enlarge button").bold = True

    p(doc, "CYP-Question", "True or false? Key Concepts is built as an accordion.")
    p(doc, "CYP - TrueFalse", "True")
    p(doc, "CYP - TrueFalse", "False")

    p(doc, "CYP-Question", "Match each Word style with what it becomes.")
    table(doc, [
        ["", "Heading 1", "", "A. A block"],
        ["", "Heading 2", "", "B. A page"],
        ["", "Heading 3", "", "C. An article"],
    ], header=False)

    doc.add_heading("Answers", 2)
    p(doc, "CYP-Question", "A")
    p(doc, "CYP-Question", "A, B, D")
    p(doc, "CYP-Question", "True")
    table(doc, [
        ["B", "Heading 1", "", "A. A block"],
        ["C", "Heading 2", "", "B. A page"],
        ["A", "Heading 3", "", "C. An article"],
    ], header=False)


def chapter_two(doc) -> None:
    doc.add_heading("Chapter 2: Interactivities", 1)
    box(doc, [
        ("Programming-Note", "Interactivity: Accordion"),
        ("Heading 4", "Learning Objectives"),
        ("Normal", "Upon completion of this chapter, learners will be able to:"),
        ("Bullet list level 1", "Ask for each interactivity by name."),
        ("Bullet list level 1", "Describe what a learner does with each one."),
    ])

    doc.add_heading("Section 2.1: Accordions and Tabs", 2)
    p(doc, "Normal", "Select each panel below to open it.")
    box(doc, [
        ("Programming-Note", "Interactivity: Accordion"),
        ("Heading 4", "Accordion 1: One Panel at a Time"),
        ("Normal", "An accordion opens one panel at a time and suits a list of "
                   "topics that are read in any order."),
        ("Heading 4", "Accordion 2: A Panel With a Figure"),
        ("Normal", "A figure inside a panel stays where the author put it, "
                   "above its caption."),
        ("image", BRIGHT, 3.0),
        ("Diagram label 1", "Figure 2-1: A Figure Inside an Accordion Panel"),
        ("Heading 4", "Accordion 3: A Panel With Bullets"),
        ("Bullet list level 1", "Panels carry bullets as well as paragraphs."),
        ("Bullet list level 1", "Each panel is one item of the component."),
    ])

    p(doc, "Normal", "Tabs suit a small number of parallel topics.")
    box(doc, [
        ("Programming-Note", "Interactivity: Horizontal Tabs"),
        ("Heading 4", "Tab 1: Horizontal"),
        ("Normal", "Horizontal tabs sit in a row above their copy."),
        ("Heading 4", "Tab 2: When to Use Them"),
        ("Normal", "Use them where the titles are short and the panels are of "
                   "similar length."),
    ])
    box(doc, [
        ("Programming-Note", "Interactivity: [Vertical] Tabs"),
        ("Heading 4", "Tab 1: Vertical"),
        ("Normal", "Vertical tabs stack down the left, which suits longer titles."),
        ("Heading 4", "Tab 2: Indication"),
        ("Normal", "A product's Indication, Administration and Safety are often "
                   "authored this way."),
        ("Heading 4", "Tab 3: Safety"),
        ("Normal", "Each tab holds one topic, so the learner can compare them."),
    ])

    doc.add_heading("Section 2.2: Hot Spots, Carousel and Popups", 2)
    p(doc, "Normal", "Select each pin on the image to learn more.")
    box(doc, [
        ("Programming-Note", "Interactivity: Hotspot Images"),
        ("Heading 4", "Hotspot 1: A Pin on the Image"),
        ("Normal", "Each pin opens its own popup over the picture."),
        ("Heading 4", "Hotspot 2: Positioning"),
        ("Normal", "Pin positions cannot be expressed in a storyboard: they are "
                   "set in base/Source and read back from there on every build."),
        ("Heading 4", "Hotspot 3: A Third Pin"),
        ("Normal", "Three pins are enough to show how the component behaves."),
        ("image", DARK, 4.5),
        ("Diagram label 1", "Figure 2-2: A Hot Spot Image"),
    ])

    p(doc, "Normal", "A carousel steps through its panels in order.")
    box(doc, [
        ("Programming-Note", "Interactivity: Narrative"),
        ("Heading 4", "First"),
        ("Normal", "A carousel panel pairs a picture with its copy."),
        ("image", BRIGHT, 3.0),
        ("Heading 4", "Second"),
        ("Normal", "The learner moves through the panels with the arrows."),
        ("image", DARK, 3.0),
        ("Heading 4", "Third"),
        ("Normal", "Use it where the order matters - a sequence or a timeline."),
        ("image", AVATAR, 1.6),
    ])

    box(doc, [
        ("Programming-Note", "Interactivity: Text with Popups"),
        ("Heading 4", "first popup"),
        ("Normal", "A popup explains a phrase without interrupting the copy."),
        ("Heading 4", "second popup"),
        ("Normal", "Each panel becomes one popup, opened from its own link."),
    ])

    doc.add_heading("Section 2.3: Sliders, Spacer and Table", 2)
    box(doc, [
        ("Programming-Note", "Interactivity: Graphic Slider"),
        ("Heading 4", "Stop 1: Before"),
        ("image", BRIGHT, 3.0),
        ("Heading 4", "Stop 2: After"),
        ("image", DARK, 3.0),
    ])
    box(doc, [
        ("Programming-Note", "Interactivity: Spacer"),
        ("Normal", "A spacer holds no copy; it separates the components around it."),
    ])
    p(doc, "Normal", "A table asked for by name becomes a component of its own:")
    p(doc, "Programming-Note", "Interactivity: Table")
    table(doc, [
        ["Component", "What the learner does"],
        ["Accordion", "Opens one panel at a time"],
        ["Tabs", "Compares parallel topics"],
        ["Hot spots", "Explores a picture"],
    ])

    doc.add_heading("Section 2.4: Video", 2)
    p(doc, "Normal",
      "A narration table becomes a video component. The line below names the "
      "file, which is also how a Key Concepts script is asked for as a video "
      "instead of an accordion - name no file and it stays an accordion.")
    p(doc, "Programming-Note", "Video: sample.mp4, Poster: poster.png")
    table(doc, [
        ["Narration for Section 2.4 Presentation", "On Screen Text"],
        ["This is the voiceover script for the video.",
         "This is what appears on screen while it plays."],
    ])

    doc.add_heading("Key Concepts (For Programming)", 2)
    table(doc, [
        ["Narration for KCs Presented by Avatar", "On Screen Text"],
        ["Accordions, tabs and hot spots each suit a different shape of content.",
         ["Choosing an Interactivity",
          ("Bullet list level 1", "Accordion: topics read in any order"),
          ("Bullet list level 1", "Tabs: parallel topics"),
          ("Bullet list level 1", "Hot spots: exploring a picture")]],
        ["Carousels, popups and sliders carry sequence, detail and comparison.",
         ["The Rest of the Set",
          ("Bullet list level 1", "Carousel: a sequence"),
          ("Bullet list level 1", "Popups: detail without interrupting"),
          ("Bullet list level 1", "Sliders: before and after")]],
    ])

    doc.add_heading("Check Your Progress", 2)
    p(doc, "CYP-Question", "Which interactivity suits a picture the learner explores?")
    p(doc, "CYP - Answer", "Tabs")
    para = doc.add_paragraph(style="CYP - Answer")
    para.add_run("Hot spots").bold = True
    p(doc, "CYP - Answer", "A carousel")
    p(doc, "CYP - Answer", "A slider")

    p(doc, "CYP-Question", "Mark whether each component is a question or not.")
    table(doc, [
        ["Question", "Not a question", ""],
        ["X", "", "Text input"],
        ["", "X", "Accordion"],
        ["X", "", "Graphic MCQ"],
    ])

    doc.add_heading("Answers", 2)
    p(doc, "CYP-Question", "B")
    p(doc, "CYP-Question", "Question, not a question, question")

    doc.add_heading("Section 2.5: Question Components", 2)
    p(doc, "Normal",
      "These three are authored as interactivities rather than as Check Your "
      "Progress questions, so they can sit anywhere in a chapter.")
    p(doc, "Interactivity-Heading-1", "Interactivity: Graphic MCQ")
    p(doc, "Interactivity-body", "Which picture is the sample figure used above?")
    p(doc, "Interactivity-Heading-2", "The dark picture.")
    para = doc.add_paragraph(style="Interactivity-Heading-2")
    para.add_run("The bright picture.").bold = True

    p(doc, "Interactivity-Heading-1", "Interactivity: Slider")
    p(doc, "Interactivity-body", "How many components does this module show?")
    p(doc, "Interactivity-Heading-2", "Scale: 1 to 20")
    p(doc, "Interactivity-Heading-2", "Answer: 18")
    p(doc, "Interactivity-Heading-2", "Labels: a few, all of them")

    p(doc, "Interactivity-Heading-1", "Interactivity: Text Input")
    p(doc, "Interactivity-body", "Which Word style becomes a page?")
    p(doc, "Interactivity-Heading-2", "Heading 1 / H1")


def back_matter(doc) -> None:
    doc.add_heading("Glossary", 1)
    table(doc, [
        ["base sample module",
         "the course in base/Source: one sample of every component the "
         "automation can build"],
        ["glossary term",
         "a word marked in the storyboard with the Glossary item character "
         "style, shown to the learner as a popup"],
        ["highlighted in boldface type",
         "the way a glossary term appears in the copy"],
    ], header=False)

    doc.add_heading("References", 1)
    p(doc, "il Reference",
      "Adapt Learning. Adapt framework documentation. Accessed 2026. "
      "https://www.adaptlearning.org/")
    p(doc, "il Reference",
      "Base Source. The sample storyboard that builds the base course. 2026.")


def build(path: Path) -> None:
    doc = Document()
    ensure_styles(doc)
    doc.core_properties.title = TITLE

    intro(doc)
    chapter_one(doc)
    chapter_two(doc)
    back_matter(doc)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    print("wrote", path)


if __name__ == "__main__":
    build(DEST / f"{TITLE}.docx")
